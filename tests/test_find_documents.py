"""find_documents.py: the search file, the scoring, and a search and --collect on made-up PDFs."""
import hashlib
import os
import shutil
import sys
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import helpers as H  # noqa: E402

FD = H.load_script("find_documents")


def pdf_bytes(pages):
    """A small valid PDF: one list of text lines per page (plain ASCII, no brackets)."""
    objs = ["<< /Type /Catalog /Pages 2 0 R >>",
            f"<< /Type /Pages /Kids [{' '.join(f'{4 + 2 * i} 0 R' for i in range(len(pages)))}] /Count {len(pages)} >>",
            "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    for i, lines in enumerate(pages):
        stream = "BT /F1 12 Tf 72 720 Td 14 TL " + " ".join(f"({x}) '" for x in lines) + " ET"
        objs.append(f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
                    f"/Resources << /Font << /F1 3 0 R >> >> /Contents {5 + 2 * i} 0 R >>")
        objs.append(f"<< /Length {len(stream)} >>\nstream\n{stream}\nendstream")
    out, offs = bytearray(b"%PDF-1.4\n"), []
    for k, o in enumerate(objs, 1):
        offs.append(len(out))
        out += f"{k} 0 obj\n{o}\nendobj\n".encode()
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode()
    out += b"".join(f"{o:010d} 00000 n \n".encode() for o in offs)
    out += f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    return bytes(out)


def title(vol, first, last):
    return ["Thailand Archaeometallurgy Project", "Small Finds Catalog", "Non Pa Wai and Non Mak La",
            f"Volume {vol}", f"Ttag #{first} - Ttag #{last}",
            "Judy Voelker and Matthew Buttacavoli with assistance from Aleshia Hospelhorn", "April 2011"]


class SearchFile(unittest.TestCase):
    def test_catalogs(self):
        s = FD.load_search("catalogs")
        self.assertEqual(len(s.targets), 18)
        self.assertEqual(s.min_pages, 3)
        self.assertTrue(s.names.search("NML Small Finds CATALOG vol2.pub"))
        t = {t.id: t for t in s.targets}["NPW NML small finds v2"]
        self.assertEqual(t.keys, ["non pa wai", "volume 2", "10028", "15985"])
        self.assertIn("buttacavoli", t.words)
        self.assertNotIn("and", t.words)
        notebook = {t.id: t for t in s.targets}["Small finds notebook v1 2010"]
        self.assertNotIn("hand", notebook.words)          # [bracketed] notes are left out

    def test_norm(self):
        self.assertEqual(FD.norm("Ttag#1791-18079, Vol. 2"), "ttag 1791 18079 volume 2")
        self.assertEqual(FD.norm("Catalogue: DataSheets"), "catalog data sheets")
        self.assertEqual(FD.norm("NML94 sherds"), "non mak la 94 sherds")


class Scoring(unittest.TestCase):
    def setUp(self):
        self.t = {t.id: t for t in FD.load_search("catalogs").targets}

    def test_volumes_apart(self):
        doc = FD.Doc("\n".join(title(2, 10028, 15985)))
        s2, _, missing2 = FD.score(self.t["NPW NML small finds v2"], doc)
        s1, _, missing1 = FD.score(self.t["NPW NML small finds v1"], doc)
        self.assertEqual(missing2, [])
        self.assertGreater(s2, 0.9)
        self.assertIn("volume 1", missing1)
        self.assertLess(s1, FD.TAKE_SCORE)

    def test_ocr_slips(self):
        doc = FD.Doc("Thailand Archaeometalurgy Projcct\nSmall Finds Catalog\nNon Pa Wai and Non Mak La\n"
                     "Volume 2\nTtag #10028 - Ttag #15985\nJudy Voe1ker and Matthew Buttacavol1\nApril 2011")
        s, _, missing = FD.score(self.t["NPW NML small finds v2"], doc)
        self.assertEqual(missing, [])
        self.assertGreater(s, FD.TAKE_SCORE)

    def test_unrelated(self):
        s, _, _ = FD.score(self.t["Phu Lon stone tools v1"], FD.Doc("Minutes of the faculty meeting, March 2004"))
        self.assertEqual(s, 0)


@unittest.skipUnless(shutil.which("pdftotext") and __import__("importlib").util.find_spec("openpyxl"),
                     "needs poppler and openpyxl")
class Run(unittest.TestCase):
    def add(self, root, db, path, content):
        p = root / "Volumes" / "ARCH1" / path
        p.parent.mkdir(parents=True, exist_ok=True)
        p.write_bytes(content)
        name = p.name
        db.execute("""INSERT INTO files(archive, unit, drive, path, name, ext, size, mtime, modified, junk, sha256, mime)
                      VALUES ('ARCH1','CumulativeArchive','THAI 2008',?,?,?,?,?,'2011-04-05 10:00:00',0,?,?)""",
                   (path, name, os.path.splitext(name)[1].lower(), len(content), p.stat().st_mtime,
                    hashlib.sha256(content).hexdigest(), "application/pdf" if name.endswith(".pdf") else ""))

    def test_search_and_collect(self):
        import openpyxl
        root, config = H.make_archive()
        db = H.connect(config)
        filler = [["Ttag 10031 bronze bracelet fragment"]] * 3
        self.add(root, db, "CumulativeArchive/THAI 2008/Catalogs/catalog v2.pdf", pdf_bytes([title(2, 10028, 15985)] + filler))
        self.add(root, db, "CumulativeArchive/THAI 2008/Catalogs/catalog v1.pdf", pdf_bytes([title(1, 5275, 9980)] + filler))
        self.add(root, db, "CumulativeArchive/THAI 2008/misc/flyer.pdf", pdf_bytes([title(3, 16011, 32385)]))
        # a control character, which Excel refuses, in a name (and so in the sheet)
        self.add(root, db, "CumulativeArchive/THAI 2008/misc/report\x01.pdf", pdf_bytes([["Annual report 2003"]] * 4))
        self.add(root, db, "CumulativeArchive/THAI 2008/misc/NML catalog.pub", b"publisher file " * 20)
        db.commit()
        H.build_unique(db)
        db.close()
        found = root / "found"
        r = H.run_script("find_documents", "catalogs", "--found-dir", str(found), "--no-ocr", config=config)
        self.assertEqual(r.returncode, 0, r.stderr)
        wb = openpyxl.load_workbook(found / "catalogs.xlsx")
        ws = wb["Matches"]
        head = [c.value for c in ws[1]]
        rows = [dict(zip(head, row)) for row in ws.iter_rows(min_row=2, values_only=True)]
        v2 = [x for x in rows if x["Target"] == "NPW NML small finds v2"]
        self.assertEqual(v2[0]["Name"], "catalog v2.pdf")
        self.assertEqual(v2[0]["Take"], "yes")
        self.assertEqual(v2[0]["Pages"], 4)
        self.assertTrue(all(x["Take"] != "yes" for x in v2[1:]))
        v1 = [x for x in rows if x["Target"] == "NPW NML small finds v1"]
        self.assertEqual((v1[0]["Name"], v1[0]["Take"]), ("catalog v1.pdf", "yes"))
        # a 1-page PDF is not read (min_pages 3), so the volume 3 flyer is not a candidate
        self.assertFalse(any(x["Name"] == "flyer.pdf" for x in rows))
        self.assertIn("NML catalog.pub", [row[3] for row in wb["By name only"].iter_rows(min_row=2, values_only=True)])
        self.assertIn("report.pdf", [row[0] for row in wb["Other long PDFs"].iter_rows(min_row=2, values_only=True)])

        # JB unmarks volume 1 and marks the Publisher file; a rerun keeps both marks
        take = head.index("Take") + 1
        for i in range(2, ws.max_row + 1):
            if ws.cell(i, 1).value == "NPW NML small finds v1" and ws.cell(i, 6).value == "catalog v1.pdf":
                ws.cell(i, take).value = ""
        bn = wb["By name only"]
        for i in range(2, bn.max_row + 1):
            if bn.cell(i, 4).value == "NML catalog.pub":
                bn.cell(i, 3).value = "x"
                pub_target = bn.cell(i, 1).value
        wb.save(found / "catalogs.xlsx")
        r = H.run_script("find_documents", "catalogs", "--found-dir", str(found), "--no-ocr", config=config)
        self.assertEqual(r.returncode, 0, r.stderr)
        self.assertIn("read before", r.stderr)               # nothing read twice
        r = H.run_script("find_documents", "catalogs", "--found-dir", str(found), "--collect", config=config)
        self.assertEqual(r.returncode, 0, r.stderr)
        got = sorted(p.name for p in (found / "catalogs").iterdir())
        self.assertTrue(pub_target.startswith("NML"), pub_target)     # NML is Non Mak La
        self.assertEqual(got, sorted(["NPW NML small finds v2 - catalog v2.pdf", f"{pub_target} - NML catalog.pub",
                                      "manifest.csv"]))
        self.assertIn("ARCH1/CumulativeArchive/THAI 2008/Catalogs/catalog v2.pdf",
                      (found / "catalogs" / "manifest.csv").read_text(encoding="utf-8-sig"))


if __name__ == "__main__":
    unittest.main()
