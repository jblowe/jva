#!/usr/bin/env python3
"""
date_report.py -- when were the archive's files created? A spreadsheet of creation dates.

    python3 scripts/date_report.py                    # writes creation_dates.xlsx in the data folder
    python3 scripts/date_report.py --by month         # one row per month instead of per year
    python3 scripts/date_report.py -o ~/Desktop/dates.xlsx

One sheet with two frequency distributions of creation dates, side by side, each with
percentages:
  * all files    every copy of every file in the catalogue (clutter such as .DS_Store, "._"
                 files and empty files left out)
  * unique files each distinct file once (identical copies together, as in the search app),
                 dated by the EARLIEST creation date among its copies
and, for both, how many files have a creation date equal to their modification date (to the
second) and how many differ, per year and in total.

Which creation date. When drive_inventory.py copied a drive onto an archive drive, each copy
got the date of copying as its creation date (most of them say 2026). The real creation
date is the one on the source drive, which drive_inventory.py recorded in that drive's CSV
(e.g. VOELKER.csv) in the data folder; this report uses that wherever it has it, and the
archive drive's own date otherwise (for folders put on the archive drives some other way).
The sheet says how many dates came from each.

Counts are plain numbers; percentages and totals are formulas, so the sheet stays correct
if you edit it. Needs openpyxl:  python3 -m pip install openpyxl
"""
import argparse
import csv
import importlib.util
import sys
import time
from collections import Counter, defaultdict
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))     # the repository, for jva
from jva.config import CONFIG  # noqa: E402
from jva.catalog import connect, log  # noqa: E402

FILE_STATUSES = ("file", "COPY ERROR")      # rows of a drive CSV that are files


def source_dates(data_dir):
    """{(archive volume, path on it): (created, modified)} from the drive CSVs, for every
    file drive_inventory.py copied. Paths are as in the catalogue."""
    out, used = {}, 0
    amap = Path(data_dir) / "archive_map.csv"
    if not amap.exists():
        return out, used
    with open(amap, newline="", encoding="utf-8-sig") as fh:
        drives = list(csv.DictReader(fh))
    for d in drives:
        name, vol, folder = d.get("archive_name"), d.get("archive_volume"), (d.get("archive_dir") or "").strip("/")
        p = Path(data_dir) / f"{name}.csv"
        if not (name and vol and folder and p.exists()):
            continue
        used += 1
        with open(p, newline="", encoding="utf-8-sig") as fh:
            for r in csv.DictReader(fh):
                if r.get("status") in FILE_STATUSES and r.get("created"):
                    out[(vol, f"{folder}/{r['path']}")] = (r["created"], r.get("modified") or "")
    return out, used


def collect(db, src, period):
    """-> (all-files table, unique-files table, counts of where dates came from). Tables are
    {period: [files, same, different]}; period is "" when there is no creation date."""
    cut = 4 if period == "year" else 7
    all_t = defaultdict(lambda: [0, 0, 0])
    first = {}                       # unique_id -> earliest creation date among its copies
    origin = Counter()
    for uid, arch, path, created, modified in db.execute(
            "SELECT unique_id, archive, path, created, modified FROM files WHERE junk = 0"):
        s = src.get((arch, path))
        if s:
            created = s[0]
            origin["source"] += 1
        else:
            origin["archive"] += 1
        created, modified = (created or "")[:19], (modified or "")[:19]
        row = all_t[created[:cut]]
        row[0] += 1
        row[1 if created and created == modified else 2] += 1
        if uid is not None and created and (uid not in first or created < first[uid]):
            first[uid] = created
    uni_t = defaultdict(lambda: [0, 0, 0])
    for uid, modified in db.execute("SELECT id, modified FROM unique_files"):
        created, modified = first.get(uid, ""), (modified or "")[:19]
        row = uni_t[created[:cut]]
        row[0] += 1
        row[1 if created and created == modified else 2] += 1
    return all_t, uni_t, origin


def periods(keys):
    """The periods that have files, in order, then "" (no date) if any. A period with no
    files in either distribution gets no row."""
    return sorted(k for k in keys if k) + ([""] if "" in keys else [])


def write_xlsx(path, all_t, uni_t, origin, period, catalogue):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter as col

    wb = Workbook()
    ws = wb.active
    ws.title = "Creation dates"
    F = "Arial"
    plain, bold = Font(name=F, size=10), Font(name=F, size=10, bold=True)
    head_fill = PatternFill("solid", fgColor="DDE4EE")
    total_fill = PatternFill("solid", fgColor="F2F2F2")
    thin = Side(style="thin", color="999999")
    INT, PCT = "#,##0", "0.0%"

    ws["A1"] = f"Creation dates of the files in the archive, by {period}"
    ws["A1"].font = Font(name=F, size=13, bold=True)
    ws["A2"] = (f"From {Path(catalogue).name}, {time.strftime('%Y-%m-%d %H:%M')}. "
                "Notes below the table.")
    ws["A2"].font = Font(name=F, size=9, italic=True)

    # two header rows: group names, then column names
    groups = (("All files (every copy)", 2), ("Unique files (identical copies counted once)", 8))
    cols = ["Files", "% of files", "Created = modified", "Created ≠ modified", "% created = modified"]
    ws.cell(5, 1, "Year created" if period == "year" else "Month created")
    for title, c0 in groups:
        ws.cell(4, c0, title)
        ws.merge_cells(start_row=4, start_column=c0, end_row=4, end_column=c0 + 4)
        for i, name in enumerate(cols):
            ws.cell(5, c0 + i, name)
    for r in (4, 5):
        for c in range(1, 13):
            if c == 7:
                continue
            cell = ws.cell(r, c)
            cell.font, cell.fill = bold, head_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = Border(bottom=thin) if r == 5 else Border()

    keys = periods(set(all_t) | set(uni_t))
    first, last = 6, 6 + len(keys) - 1
    total = last + 1
    for i, k in enumerate(keys):
        r = first + i
        ws.cell(r, 1, k or "No date")                     # text, so years aren't summed or comma'd
        for t, c0 in ((all_t, 2), (uni_t, 8)):
            n, same, diff = t.get(k, (0, 0, 0))
            b, d = col(c0), col(c0 + 2)
            ws.cell(r, c0, n)
            ws.cell(r, c0 + 1, f"=IF({b}${total}=0,0,{b}{r}/{b}${total})")
            ws.cell(r, c0 + 2, same)
            ws.cell(r, c0 + 3, diff)
            ws.cell(r, c0 + 4, f'=IF({b}{r}=0,"",{d}{r}/{b}{r})')
    ws.cell(total, 1, "Total")
    for c0 in (2, 8):
        b, d = col(c0), col(c0 + 2)
        for j in (0, 2, 3):
            L = col(c0 + j)
            ws.cell(total, c0 + j, f"=SUM({L}{first}:{L}{last})")
        ws.cell(total, c0 + 1, f"=IF({b}{total}=0,0,{b}{total}/{b}{total})")
        ws.cell(total, c0 + 4, f'=IF({b}{total}=0,"",{d}{total}/{b}{total})')
    for r in range(first, total + 1):
        ws.cell(r, 1).font = bold if r == total else plain
        ws.cell(r, 1).alignment = Alignment(horizontal="left")
        for c in list(range(2, 7)) + list(range(8, 13)):
            cell = ws.cell(r, c)
            cell.font = bold if r == total else plain
            cell.number_format = PCT if (c - 2) % 6 in (1, 4) else INT
            if r == total:
                cell.fill, cell.border = total_fill, Border(top=thin)
    ws.cell(total, 1).fill, ws.cell(total, 1).border = total_fill, Border(top=thin)

    # created vs modified, both distributions
    r0 = total + 3
    ws.cell(r0, 1, "Creation date vs modification date").font = Font(name=F, size=11, bold=True)
    for j, h in enumerate(["", "All files", "% of all files", "Unique files", "% of unique files"]):
        c = ws.cell(r0 + 1, 1 + j, h)
        c.font, c.fill, c.border = bold, head_fill, Border(bottom=thin)
        c.alignment = Alignment(horizontal="center", wrap_text=True)
    rows = (("Created = modified (to the second)", 4), ("Created ≠ modified", 5))
    for i, (label, off) in enumerate(rows):
        r = r0 + 2 + i
        ws.cell(r, 1, label).font = plain
        for j, c0 in enumerate((2, 8)):
            ws.cell(r, 2 + 2 * j, f"={col(c0 + off - 2)}{total}")
            ws.cell(r, 3 + 2 * j, f"=IF({col(c0)}{total}=0,0,{col(2 + 2 * j)}{r}/{col(c0)}{total})")
    r = r0 + 4
    ws.cell(r, 1, "Total").font = bold
    for j in (0, 1):
        L = col(2 + 2 * j)
        ws.cell(r, 2 + 2 * j, f"=SUM({L}{r0 + 2}:{L}{r0 + 3})")
        ws.cell(r, 3 + 2 * j, f"=SUM({col(3 + 2 * j)}{r0 + 2}:{col(3 + 2 * j)}{r0 + 3})")
    for rr in range(r0 + 2, r0 + 5):
        for c in range(2, 6):
            cell = ws.cell(rr, c)
            cell.font = bold if rr == r else plain
            cell.number_format = PCT if c in (3, 5) else INT
            if rr == r:
                cell.fill, cell.border = total_fill, Border(top=thin)
    ws.cell(r, 1).fill, ws.cell(r, 1).border = total_fill, Border(top=thin)

    notes = [
        "Notes",
        "All files: every copy of every file in the catalogue, leaving out clutter (.DS_Store, "
        "\"._\" files, empty files). Unique files: identical copies counted once.",
        "Creation date: for files that drive_inventory.py copied from a source drive, the date on "
        "the source drive, from that drive's CSV in the data folder "
        f"({origin['source']:,} files). Their copies on the archive drive are dated by the copying "
        f"(mostly 2026), so that date is not used. Other files: the archive drive's own date "
        f"({origin['archive']:,} files).",
        "A unique file's creation date is the earliest among its copies; its modification date is "
        "the earliest modification date among its copies (as in the search app).",
        "Created = modified: the two dates are identical to the second. Copying a file to a new "
        "disk often resets its creation date but not its modification date, so 'created after "
        "modified' is common in old material.",
        "Dates in 1969 or 1970 are almost always a missing date (the zero of a Mac or Unix clock), "
        "and late 1979 or 1 Jan 1980 usually the zero of a DOS/FAT drive, not real dates; dates "
        "in the future come from wrong clocks.",
        "Counts are numbers; percentages and totals are formulas. Rerun: python3 scripts/date_report.py",
    ]
    rn = r + 3
    for i, t in enumerate(notes):
        c = ws.cell(rn + i, 1, t)
        c.font = bold if i == 0 else plain
        c.alignment = Alignment(wrap_text=False)

    ws.column_dimensions["A"].width = 34
    for c in list(range(2, 7)) + list(range(8, 13)):
        ws.column_dimensions[col(c)].width = 13
    ws.column_dimensions["G"].width = 3
    ws.row_dimensions[5].height = 30
    ws.freeze_panes = "B6"
    wb.save(path)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--by", choices=("year", "month"), default="year", help="one row per year (default) or month")
    ap.add_argument("-o", "--output", help="the spreadsheet (default: creation_dates.xlsx in the data folder)")
    ap.add_argument("--catalogue", default=str(CONFIG.catalogue))
    a = ap.parse_args()
    if importlib.util.find_spec("openpyxl") is None:
        sys.exit("This needs openpyxl:  python3 -m pip install openpyxl")
    out = Path(a.output).expanduser() if a.output else CONFIG.data_dir / "creation_dates.xlsx"
    t0 = time.time()
    src, n_drives = source_dates(CONFIG.data_dir)
    log(f"source-drive dates for {len(src):,} files from {n_drives} drive CSVs")
    db = connect(a.catalogue, readonly=True)
    all_t, uni_t, origin = collect(db, src, a.by)
    db.close()
    write_xlsx(out, all_t, uni_t, origin, a.by, a.catalogue)
    n_all, n_uni = sum(v[0] for v in all_t.values()), sum(v[0] for v in uni_t.values())
    print(f"Wrote {out}: {n_all:,} files, {n_uni:,} unique files ({time.time() - t0:.0f}s)")


if __name__ == "__main__":
    main()
