#!/usr/bin/env python3
"""
word_report.py -- the archive's Microsoft Word documents and how many pages each has.

    python3 scripts/word_report.py                  # writes word_documents.xlsx in the data folder
    python3 scripts/word_report.py -o ~/Desktop/word.xlsx
    python3 scripts/word_report.py --max-minutes 20 # stop reading after 20 minutes; rerun to finish

Counts each distinct Word document once (identical copies together, as in the search app):
.doc, .docx, .dot, .dotx and .docm files, and old Mac Word files without an extension
(recognised by their Mac type and creator codes). Office lock files (~$...) are left out.

Pages. Word stores the page count it last worked out when it saved a document:
  .docx etc.   docProps/app.xml inside the file ("Pages")
  .doc etc.    the document's summary information (Word 97 and later, and Word 6/95)
Old Mac Word files (Word 5 and earlier) have no stored count; for those, and whenever the
stored count is missing or 0, the page count of the PDF copy that make_derivatives.py made
is used, if there is one. A stored count can also be plainly wrong (a 16-page essay saved as
"1 page", by a program that never paginated it): when there is no count at all, or more than
1,000 words per page, the document is paginated again by LibreOffice (converted to PDF in a
temporary folder), if it is installed. The "Pages from" column says which. Word's own count is the count
at the last save, with that computer's fonts and paper size, so it can differ slightly from
what Word shows today.

Reads one copy of each document from whichever archive drive is plugged in (read only).
Results are kept in word_pages.json in the data folder, so a rerun only reads documents it
hasn't seen. Needs openpyxl and olefile:  python3 -m pip install openpyxl olefile
"""
import argparse
import importlib.util
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import time
import zipfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))     # the repository, for jva
from jva.config import CONFIG  # noqa: E402
from jva.catalog import connect, has_table, log  # noqa: E402

WORD_EXT = (".doc", ".docx", ".dot", ".dotx", ".docm", ".dotm")
WORD_MAC_TYPES = ("WDBN", "W8BN", "W6BN", "WXBN", "W8TN", "WTMP", "W6TN", "WDTN")
SELECT = f"""
    SELECT u.id, u.sha256, u.name, u.ext, u.format, u.mac_type, u.size_bytes, u.created, u.modified,
           u.copies, u.archive, u.path, COALESCE(u.category, 'Unclassified')
    FROM unique_files u
    WHERE (lower(u.ext) IN ({', '.join(repr(e) for e in WORD_EXT)})
           OR (u.mac_creator = 'MSWD' AND u.mac_type IN ({', '.join(repr(t) for t in WORD_MAC_TYPES)})))
      AND substr(u.name, 1, 2) <> '~$'
    ORDER BY u.name COLLATE NOCASE"""


# ---------- reading page counts ----------

def pages_docx(path):
    with zipfile.ZipFile(path) as z:
        xml = z.read("docProps/app.xml").decode("utf-8", "replace")
    pages = re.search(r"<Pages>(\d+)</Pages>", xml)
    words = re.search(r"<Words>(\d+)</Words>", xml)
    return (int(pages.group(1)) if pages else None), (int(words.group(1)) if words else None)


def pages_ole(path):
    import olefile
    with olefile.OleFileIO(str(path)) as ole:
        m = ole.get_metadata()
        return m.num_pages, m.num_words


def pages_pdf(path):
    data = Path(path).read_bytes()
    if importlib.util.find_spec("pypdf"):
        try:
            import io
            import pypdf
            return len(pypdf.PdfReader(io.BytesIO(data)).pages)
        except Exception:
            pass
    n = len(re.findall(rb"/Type\s*/Page(?![a-z])", data))
    return n or None


def read_document(path):
    """-> (pages, words, how) from the file itself; pages None if it has no usable count."""
    with open(path, "rb") as fh:
        head = fh.read(8)
    try:
        if head.startswith(b"PK"):
            p, w = pages_docx(path)
            return p, w, "Word's saved count (.docx)"
        if head.startswith(b"\xd0\xcf\x11\xe0"):
            p, w = pages_ole(path)
            return p, w, "Word's saved count"
    except Exception as e:                       # damaged or not really a Word file
        return None, None, f"unreadable ({type(e).__name__})"
    return None, None, "no stored count (old Mac Word)"


def a_copy(db, uid, volumes_root):
    for archive, path, size in db.execute("SELECT archive, path, size FROM files WHERE unique_id=? "
                                          "ORDER BY archive, path", (uid,)):
        p = Path(volumes_root) / archive / path
        try:
            if p.is_file() and p.stat().st_size == size:
                return p
        except OSError:
            continue
    return None


def pdf_copy(db, sha, roots):
    if not sha or not has_table(db, "derivatives"):
        return None
    cols = {r[1] for r in db.execute("PRAGMA table_info(derivatives)")}
    if "pdf" not in cols:
        return None
    r = db.execute("SELECT root, pdf FROM derivatives WHERE sha256=? AND pdf IS NOT NULL", (sha,)).fetchone()
    if not r:
        return None
    for root in list(roots) + [r[0]]:
        if root and (Path(root) / r[1]).is_file():
            return Path(root) / r[1]
    return None


# A saved count is sometimes plainly wrong: a 16-page essay recorded as 1 page, because the
# program that last saved it never paginated it. More than this many words per page means the
# count can't be right, and the document is paginated again with LibreOffice.
MAX_WORDS_PER_PAGE = 1000
BATCH = 20


def doubtful(res):
    if not res or res.get("how", "").startswith("counted by LibreOffice"):     # tried already
        return False
    if not res.get("pages"):
        return True
    return bool(res.get("words")) and res["words"] / res["pages"] > MAX_WORDS_PER_PAGE


def find_soffice():
    return shutil.which("soffice") or next(
        (p for p in ("/Applications/LibreOffice.app/Contents/MacOS/soffice",
                     os.path.expanduser("~/Applications/LibreOffice.app/Contents/MacOS/soffice"))
         if os.path.exists(p)), None)


def recount(db, cache, docs, volumes_root, deadline):
    """Paginate the doubtful documents with LibreOffice (converted to PDF, BATCH at a time)."""
    soffice = find_soffice()
    todo = [(d[0], d[1] or f"u{d[0]}") for d in docs if doubtful(cache.get(d[1] or f"u{d[0]}"))]
    if not todo:
        return
    if not soffice:
        log(f"{len(todo):,} documents have no believable saved page count; install LibreOffice to count them")
        return
    log(f"{len(todo):,} documents have no believable saved page count: counting their pages with LibreOffice")
    profile = Path(tempfile.mkdtemp(prefix="jva-lo-"))
    done = 0
    for i in range(0, len(todo), BATCH):
        if time.time() > deadline:
            log("time's up: rerun to count the rest")
            break
        with tempfile.TemporaryDirectory() as td:
            batch = {}
            for n, (uid, key) in enumerate(todo[i:i + BATCH]):
                src = a_copy(db, uid, volumes_root)
                if src is None:
                    continue
                dst = Path(td) / f"d{n}{src.suffix.lower() if src.suffix.lower() in WORD_EXT else '.doc'}"
                shutil.copyfile(src, dst)
                batch[dst] = key
            if not batch:
                continue
            try:
                subprocess.run([soffice, f"-env:UserInstallation={profile.as_uri()}", "--headless",
                                "--convert-to", "pdf", "--outdir", td, *map(str, batch)],
                               capture_output=True, timeout=60 + 30 * len(batch))
            except subprocess.TimeoutExpired:
                pass
            for dst, key in batch.items():
                pdf = dst.with_suffix(".pdf")
                if pdf.exists():
                    cache[key].update(saved=cache[key].get("pages"), pages=pages_pdf(pdf),
                                      how="counted by LibreOffice")
                else:                                  # keep Word's count, doubtful or not
                    cache[key]["how"] = "counted by LibreOffice: couldn't open it"
                done += 1
        log(f"  counted {done:,} of {len(todo):,}")
    shutil.rmtree(profile, ignore_errors=True)


def collect(db, cache, volumes_root, roots, deadline):
    docs = db.execute(SELECT).fetchall()
    todo = [d for d in docs if (d[1] or f"u{d[0]}") not in cache]
    log(f"{len(docs):,} distinct Word documents; {len(todo):,} to read ({len(docs) - len(todo):,} from earlier runs)")
    done = unplugged = 0
    last = time.time()
    for uid, sha, *_ in todo:
        if time.time() > deadline:
            log("time's up: rerun to read the rest")
            break
        key = sha or f"u{uid}"
        path = a_copy(db, uid, volumes_root)
        if path is None:
            unplugged += 1
            continue                             # not cached: tried again next time
        pages, words, how = read_document(path)
        if not pages:
            pdf = pdf_copy(db, sha, roots)
            if pdf is not None:
                try:
                    pages, how = pages_pdf(pdf), "PDF copy"
                except OSError:
                    pass
        cache[key] = {"pages": pages or None, "words": words or None, "how": how}
        done += 1
        if time.time() - last > 30:
            last = time.time()
            log(f"  read {done:,} of {len(todo):,}")
    if unplugged:
        log(f"{unplugged:,} documents are only on drives that aren't plugged in; plug them in and rerun")
    return docs


# ---------- the spreadsheet ----------

def write_xlsx(path, docs, cache):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter as col

    F = "Arial"
    plain, bold = Font(name=F, size=10), Font(name=F, size=10, bold=True)
    head_fill = PatternFill("solid", fgColor="DDE4EE")
    thin = Side(style="thin", color="999999")
    wb = Workbook()
    summ = wb.active
    summ.title = "Summary"
    ws = wb.create_sheet("Documents")

    heads = ["Name", "Where (archive drive: path)", "Category", "Word version", "Size (bytes)", "Modified",
             "Copies", "Pages", "Words", "Pages from", "Word's saved count, if different"]
    widths = [40, 80, 18, 30, 12, 18, 8, 8, 10, 30, 14]
    for i, (h, w) in enumerate(zip(heads, widths), 1):
        c = ws.cell(1, i, h)
        c.font, c.fill, c.border = bold, head_fill, Border(bottom=thin)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[col(i)].width = w
    for r, (uid, sha, name, ext, fmt, mtype, size, created, modified, copies, archive, rel, cat) in enumerate(docs, 2):
        res = cache.get(sha or f"u{uid}")
        version = fmt or {"WDBN": "Microsoft Word for Macintosh", "W8BN": "Microsoft Word 98 (Mac)",
                          "W6BN": "Microsoft Word 6 (Mac)", "WXBN": "Microsoft Word 2001/X (Mac)"}.get(mtype, "")
        vals = [name, f"{archive}: {rel}", cat, version, size, (modified or "")[:16], copies,
                res["pages"] if res else None, res["words"] if res else None,
                (res["how"] if res else "drive not plugged in"),
                (res.get("saved") if res and res.get("saved") != res.get("pages") else None)]
        for i, v in enumerate(vals, 1):
            c = ws.cell(r, i, v)
            c.font = plain
            if i in (5, 7, 8, 9, 11):
                c.number_format = "#,##0"
    last = len(docs) + 1
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:K{last}"

    # summary: formulas over the Documents sheet
    D = "Documents"
    rng = lambda c: f"{D}!${c}$2:${c}${last}"           # noqa: E731
    summ.column_dimensions["A"].width = 46
    for c in "BCD":
        summ.column_dimensions[c].width = 14
    summ["A1"] = "Microsoft Word documents in the archive"
    summ["A1"].font = Font(name=F, size=13, bold=True)
    summ["A2"] = (f"Each distinct document once (identical copies counted once), "
                  f"{time.strftime('%Y-%m-%d %H:%M')}. Details on the Documents sheet.")
    summ["A2"].font = Font(name=F, size=9, italic=True)

    def table(r, title, rows, cols):
        summ.cell(r, 1, title).font = Font(name=F, size=11, bold=True)
        for j, h in enumerate(cols, 1):
            c = summ.cell(r + 1, j, h)
            c.font, c.fill, c.border = bold, head_fill, Border(bottom=thin)
        for i, vals in enumerate(rows):
            for j, v in enumerate(vals, 1):
                c = summ.cell(r + 2 + i, j, v)
                c.font = plain
                c.number_format = "#,##0" if j > 1 else "General"
        return r + 2 + len(rows) + 2

    r = table(4, "Totals", [
        ("Distinct Word documents", f"=COUNTA({rng('A')})"),
        ("with a page count", f"=COUNT({rng('H')})"),
        ("without one", "=B6-B7"),
        ("Total pages (of those with a count)", f"=SUM({rng('H')})"),
        ("Average pages per document", "=IF(B7=0,0,B9/B7)"),
        ("Median pages", f"=IF(B7=0,0,MEDIAN({rng('H')}))"),
        ("Total words (where Word recorded them)", f"=SUM({rng('I')})"),
    ], ["", "Number"])
    summ["B10"].number_format = summ["B11"].number_format = "#,##0.0"

    bands = [(1, 1), (2, 2), (3, 5), (6, 10), (11, 25), (26, 50), (51, 100), (101, 250), (251, 10 ** 6)]
    rows = []
    for lo, hi in bands:
        label = (f"{lo} page" + ("" if lo == 1 else "s")) if lo == hi else \
                (f"{lo} or more pages" if hi >= 10 ** 6 else f"{lo}–{hi} pages")
        rows.append((label, f'=COUNTIFS({rng("H")},">={lo}",{rng("H")},"<={hi}")',
                     f'=SUMIFS({rng("H")},{rng("H")},">={lo}",{rng("H")},"<={hi}")'))
    rows.append(("No page count", "=B6-B7", 0))
    start = r
    r = table(r, "Documents by length", rows, ["Length", "Documents", "Pages"])
    for rr in range(start + 2, r - 2):
        summ.cell(rr, 4, f"=IF($B$6=0,0,B{rr}/$B$6)").number_format = "0.0%"
    summ.cell(start + 1, 4, "% of documents").font = bold
    summ.cell(start + 1, 4).fill = head_fill

    def by(column, values, title, r):
        rows = [(v, f'=COUNTIF({rng(column)},"{v}")', f'=SUMIF({rng(column)},"{v}",{rng("H")})') for v in values]
        return table(r, title, rows, [title.split(" by ")[-1].capitalize(), "Documents", "Pages"])

    cats = sorted({d[12] for d in docs})
    r = by("C", cats, "Documents by category", r)
    versions = sorted({ws.cell(i, 4).value or "" for i in range(2, last + 1)} - {""})
    r = by("D", versions, "Documents by Word version", r)
    hows = sorted({ws.cell(i, 10).value or "" for i in range(2, last + 1)} - {""})
    r = by("J", hows, "Documents by where the page count came from", r)

    notes = ["Notes",
             "Pages: the count Word saved inside the document (docProps/app.xml in .docx files; the "
             "summary information in .doc files). Old Mac Word files (Word 5 and earlier) have none: for "
             "those, and when the saved count is 0 or missing, the page count of the PDF copy made by "
             "make_derivatives.py is used. When there is no count, or more than 1,000 words per page "
             "(a count that can't be right), LibreOffice paginates the document again; the Documents "
             "sheet then shows Word's saved count beside it.",
             "Word's saved count is from the last time the document was saved, with that computer's "
             "fonts and paper size; opening it today can give a slightly different number.",
             "Rerun: python3 scripts/word_report.py"]
    for i, t in enumerate(notes):
        summ.cell(r + i, 1, t).font = bold if i == 0 else plain
    wb.save(path)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("-o", "--output", help="the spreadsheet (default: word_documents.xlsx in the data folder)")
    ap.add_argument("--catalogue", default=str(CONFIG.catalogue))
    ap.add_argument("--max-minutes", type=float, default=0, help="stop reading after this long (0: no limit)")
    a = ap.parse_args()
    missing = [m for m in ("openpyxl", "olefile") if importlib.util.find_spec(m) is None]
    if missing:
        sys.exit(f"This needs {' and '.join(missing)}:  python3 -m pip install {' '.join(missing)}")
    out = Path(a.output).expanduser() if a.output else CONFIG.data_dir / "word_documents.xlsx"
    cache_file = CONFIG.data_dir / "word_pages.json"
    cache = json.loads(cache_file.read_text()) if cache_file.exists() else {}
    deadline = time.time() + a.max_minutes * 60 if a.max_minutes else float("inf")
    db = connect(a.catalogue, readonly=True)
    roots = CONFIG.app.get("derivatives_roots") or []
    if CONFIG.derivatives_dir:
        roots = [str(CONFIG.derivatives_dir)] + list(roots)
    try:
        docs = collect(db, cache, CONFIG.volumes_root, roots, deadline)
        recount(db, cache, docs, CONFIG.volumes_root, deadline)
    finally:
        cache_file.write_text(json.dumps(cache))
    write_xlsx(out, docs, cache)
    db.close()
    have = sum(1 for d in docs if (cache.get(d[1] or f"u{d[0]}") or {}).get("pages"))
    print(f"Wrote {out}: {len(docs):,} distinct Word documents, {have:,} with a page count")


if __name__ == "__main__":
    main()
