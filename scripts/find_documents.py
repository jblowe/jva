#!/usr/bin/env python3
"""
find_documents.py -- find particular documents in the archive by what their title pages say,
and gather copies of them in a folder outside the archive.

    python3 scripts/find_documents.py catalogs             # search: writes ~/jva_found/catalogs.xlsx
    python3 scripts/find_documents.py catalogs --collect   # copy the files marked Take to ~/jva_found/catalogs/
    python3 scripts/find_documents.py --list               # the searches in rules/searches/
    python3 scripts/find_documents.py ~/my_search.csv      # a search file kept somewhere else
    python3 scripts/find_documents.py catalogs --max-minutes 30   # stop reading after 30 minutes; rerun to finish

A search is a CSV file, normally rules/searches/<name>.csv, with one row per document wanted:
  id      a short name for it, used in the spreadsheet and in the copies' file names
  title   what its title page says (several lines are fine; [bracketed] notes are ignored)
  keys    words or numbers that tell it apart from its neighbours, ";"-separated: a volume
          number, the Ttag range, a year ("volume 2; 10028; 15985")
  notes   anything, for people
Lines at the top starting with "#" are comments, and can set the search's options:
  # names: catalog|data ?sheet     files whose names match this (a regular expression, any
                                   case) are listed whatever their kind, and are read
                                   whatever their length
  # min_pages: 3                   shorter PDFs are not read (default 2)

What is read. Each distinct file once (identical copies together, as in the search app), from
whichever archive drive holding a copy is plugged in, read only:
  PDFs         the page count (pdfinfo) and the text of pages 1-2 (pdftotext); a PDF with no
               text (a scan) has page 1 read by OCR (Tesseract), and page 2 if page 1 is blank
  Word, RTF    the start of the text (macOS textutil); old Mac Word files textutil can't read
               use the PDF copy make_derivatives.py made, if there is one
  the rest     (Publisher, InDesign, Pages, ...) only by name and folder, and only if the name
               matches the search's "names"
What was read is kept in document_text.sqlite in the data folder (not in the catalogue), so a
rerun, or a different search, reads only files it hasn't read before.

Scoring. Each file's text, plus its name and the names of its last two folders, is compared
with each title: the share of the title's words it contains (numbers count double; words of 5
letters or more may differ by a typo or an OCR slip) and the share of the keys it contains.
Score = 0.6 x words + 0.4 x keys, from 0 to 1.

The spreadsheet (<found folder>/<search>.xlsx; the found folder is ~/jva_found unless
"found_dir" in config.json or --found-dir says otherwise):
  Targets          the documents wanted, and how many candidates each has
  Matches          each one's best candidates (--top, default 5), best first. Take is "yes" where
                   every key was found and the score is 0.75 or more. Check them; mark or unmark
                   Take yourself (yes / x / blank): a rerun keeps your marks
  By name only     files whose names match "names" but whose text couldn't be read
  Other long PDFs  PDFs of min_pages or more that matched nothing, with the start of page 1
--collect copies every file marked Take into <found folder>/<search>/ as "<id> - <name>", with
manifest.csv saying where each came from. Nothing in the archive is changed.

Needs openpyxl; poppler (pdfinfo, pdftotext, pdftoppm) and tesseract: brew install poppler tesseract
"""
import argparse
import concurrent.futures as cf
import csv
import difflib
import importlib.util
import io
import re
import shutil
import sqlite3
import subprocess
import sys
import tempfile
import time
import unicodedata
import zipfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))     # the repository, for jva
from jva.config import CONFIG, REPO  # noqa: E402
from jva.catalog import connect, has_table, log  # noqa: E402

SEARCH_DIR = REPO / "rules" / "searches"
TEXT_EXT = {".doc", ".docx", ".dot", ".dotx", ".docm", ".rtf", ".odt"}
WORD_MAC = {"WDBN", "W8BN", "W6BN", "WXBN", "W8TN", "WTMP", "W6TN", "WDTN"}
TEXT_LIMIT = 4000          # characters of a file's text kept
MIN_SCORE = 0.4            # lower scores are not listed
TAKE_SCORE = 0.75          # Take is "yes" at this score with every key found
ENOUGH = 40                # letters and digits a page needs to count as having text


# ---------- the search file ----------

class Target:
    def __init__(self, id, title, keys="", notes=""):
        self.id, self.title, self.notes = id.strip(), (title or "").strip(), (notes or "").strip()
        words = norm(re.sub(r"\[[^\]]*\]", " ", self.title)).split()
        self.words = [w for w in dict.fromkeys(words) if w not in STOP and (len(w) > 1 or w.isdigit())]
        self.keys = [k for k in (norm(k) for k in (keys or "").split(";")) if k]

    def one_line(self):
        return " / ".join(line.strip() for line in self.title.splitlines() if line.strip())


class Search:
    def __init__(self, name, path, targets, names="", min_pages=2):
        self.name, self.path, self.targets = name, path, targets
        self.names = re.compile(names, re.I) if names else None
        self.min_pages = min_pages


def load_search(spec):
    p = Path(spec).expanduser()
    if not p.suffix:
        p = SEARCH_DIR / f"{spec}.csv"
    if not p.is_file():
        sys.exit(f"No search file {p}  (the searches: python3 scripts/find_documents.py --list)")
    lines = p.read_text(encoding="utf-8-sig").splitlines(keepends=True)
    opts, i = {}, 0
    while i < len(lines) and (lines[i].startswith("#") or not lines[i].strip()):
        m = re.match(r"#\s*(names|min_pages)\s*:\s*(.*?)\s*$", lines[i], re.I)
        if m:
            opts[m.group(1).lower()] = m.group(2)
        i += 1
    targets = [Target(r["id"], r.get("title"), r.get("keys"), r.get("notes"))
               for r in csv.DictReader(io.StringIO("".join(lines[i:]))) if (r.get("id") or "").strip()]
    ids = [t.id for t in targets]
    if len(set(ids)) != len(ids):
        sys.exit(f"{p}: the ids must all be different")
    return Search(p.stem, p, targets, opts.get("names", ""), int(opts.get("min_pages") or 2))


# ---------- comparing text ----------

STOP = {"and", "of", "the", "at", "with", "by", "from", "for", "in", "a", "an", "also", "on", "to"}
SAME = [(r"\bvol\b", "volume"), (r"\bcatalogue(s?)\b", r"catalog\1"), (r"\bdatasheets?\b", "data sheets"),
        (r"\bdata sheet\b", "data sheets"),
        # site abbreviations, common in file and folder names ("PL" is left alone: too common)
        (r"\bnml\b", "non mak la"), (r"\bnpw\b", "non pa wai"), (r"\bnkh\b", "nil kham haeng")]


def norm(text):
    """Lower case, accents off, only letters and digits, letters and digits apart
    ("Ttag#1791-18079" -> "ttag 1791 18079"), and a few spellings made the same ("vol" -> "volume",
    "NML" -> "non mak la")."""
    t = unicodedata.normalize("NFKD", text or "").encode("ascii", "ignore").decode().lower()
    t = re.sub(r"(?<=[a-z])(?=[0-9])|(?<=[0-9])(?=[a-z])", " ", t)
    t = re.sub(r"[^a-z0-9]+", " ", t)
    for a, b in SAME:
        t = re.sub(a, b, t)
    return " ".join(t.split())


class Doc:
    def __init__(self, text):
        self.text = norm(text)
        self.padded = f" {self.text} "
        self.words = set(self.text.split())
        self.by_len = {}
        for w in self.words:
            self.by_len.setdefault(len(w), []).append(w)

    def has(self, w):
        """The word, or (5 letters or more, not a number) one differing by a slip."""
        if w in self.words:
            return True
        if len(w) < 5 or w.isdigit():
            return False
        near = [x for n in range(len(w) - 2, len(w) + 3) for x in self.by_len.get(n, ())]
        return bool(difflib.get_close_matches(w, near, n=1, cutoff=0.8))

    def has_key(self, k):
        if f" {k} " in self.padded:
            return True
        # a phrase with no numbers may have slips in it ("archaeometalurgy project")
        words = k.split()
        return len(words) > 1 and not any(w.isdigit() for w in words) and all(self.has(w) for w in words)


def weight(w):
    return 2 if w.isdigit() else 1


def score(target, doc):
    """-> (score 0-1, keys found, keys missing)."""
    total = sum(weight(w) for w in target.words) or 1
    exact = sum(weight(w) for w in target.words if w in doc.words)
    if exact / total < 0.25 and not any(f" {k} " in doc.padded for k in target.keys):
        return 0.0, [], target.keys                  # nowhere near: skip the slow part
    words = sum(weight(w) for w in target.words if doc.has(w)) / total
    found = [k for k in target.keys if doc.has_key(k)]
    keys = len(found) / len(target.keys) if target.keys else words
    return round(0.6 * words + 0.4 * keys, 3), found, [k for k in target.keys if k not in found]


# ---------- reading the files ----------

def run(cmd, timeout):
    return subprocess.run(cmd, capture_output=True, timeout=timeout)


def enough(text):
    return len(re.findall(r"[A-Za-z0-9]", text or "")) >= ENOUGH


def pdf_pages(path):
    m = re.search(rb"^Pages:\s+(\d+)", run(["pdfinfo", str(path)], 60).stdout, re.M)
    return int(m.group(1)) if m else None


def pdf_text(path, first=1, last=2):
    r = run(["pdftotext", "-q", "-f", str(first), "-l", str(last), "-enc", "UTF-8", str(path), "-"], 120)
    return r.stdout.decode("utf-8", "replace")


def ocr_page(path, page):
    with tempfile.TemporaryDirectory() as td:
        run(["pdftoppm", "-q", "-f", str(page), "-l", str(page), "-r", "200", "-gray", "-png",
             str(path), f"{td}/p"], 180)
        images = sorted(Path(td).glob("p*.png"))
        if not images:
            return ""
        return run(["tesseract", str(images[0]), "stdout", "--psm", "3"], 180).stdout.decode("utf-8", "replace")


def read_pdf(path, min_pages, always, ocr):
    """-> {pages, text, how}. text None: not read yet (too short, or waiting for OCR)."""
    pages = pdf_pages(path)
    if pages is not None and pages < min_pages and not always:
        return {"pages": pages, "text": None, "how": f"not read: {pages} page{'s' if pages != 1 else ''}"}
    text = pdf_text(path)
    if enough(text):
        return {"pages": pages, "text": text, "how": "text"}
    if not ocr:
        return {"pages": pages, "text": None, "how": "no text (a scan): not OCR'd"}
    o = ocr_page(path, 1)
    if not enough(o) and (pages or 2) >= 2:
        o = (o + "\n" + ocr_page(path, 2)).strip()
    return {"pages": pages, "text": (text + "\n" + o).strip(), "how": "OCR" if enough(o) else "OCR: nothing legible"}


def read_word(path, pdf_copy):
    """The start of a Word/RTF document's text."""
    if shutil.which("textutil"):
        with tempfile.TemporaryDirectory() as td:
            src = path
            if Path(path).suffix.lower() not in TEXT_EXT:      # old Mac files without an extension
                src = Path(td) / "d.doc"
                shutil.copyfile(path, src)
            try:
                t = run(["textutil", "-convert", "txt", "-stdout", str(src)], 60).stdout.decode("utf-8", "replace")
            except subprocess.TimeoutExpired:
                t = ""
        if enough(t):
            return {"pages": None, "text": t, "how": "textutil"}
    try:
        with zipfile.ZipFile(path) as z:                         # .docx without textutil
            xml = z.read("word/document.xml").decode("utf-8", "replace")
        t = re.sub(r"<[^>]+>", " ", xml.replace("</w:p>", "\n"))
        if enough(t):
            return {"pages": None, "text": t, "how": "docx"}
    except Exception:
        pass
    if pdf_copy:
        t = pdf_text(pdf_copy)
        if enough(t):
            return {"pages": None, "text": t, "how": "PDF copy"}
    return {"pages": None, "text": "", "how": "no text found"}


def a_copy(db, uid, volumes_root):
    for archive, path, size in db.execute("SELECT archive, path, size FROM files WHERE unique_id=? "
                                          "ORDER BY archive, path", (uid,)):
        p = Path(volumes_root) / archive / path
        try:
            if p.is_file() and p.stat().st_size == size:
                return p
        except OSError:
            continue
    return None


def pdf_copies(db, roots):
    """sha256 -> the PDF copy make_derivatives.py made of an old document."""
    if not has_table(db, "derivatives") or "pdf" not in {r[1] for r in db.execute("PRAGMA table_info(derivatives)")}:
        return {}
    out = {}
    for sha, root, rel in db.execute("SELECT sha256, root, pdf FROM derivatives WHERE pdf IS NOT NULL AND pdf <> ''"):
        for r in list(roots) + [root]:
            if r and (Path(r) / rel).is_file():
                out[sha] = Path(r) / rel
                break
    return out


# ---------- the cache of what was read ----------

def open_cache(path):
    c = sqlite3.connect(path)
    c.execute("CREATE TABLE IF NOT EXISTS doc_text(key TEXT PRIMARY KEY, pages INTEGER, text TEXT, "
              "how TEXT, read_at TEXT)")
    return c


def cached(cache):
    return {k: {"pages": p, "text": t, "how": h} for k, p, t, h in cache.execute("SELECT key, pages, text, how FROM doc_text")}


# ---------- the files ----------

FILE_COLS = ("id", "sha256", "name", "ext", "mime", "mac_type", "mac_creator", "size_bytes", "modified",
             "copies", "archive", "path", "category")


class File:
    def __init__(self, row):
        for k, v in zip(FILE_COLS, row):
            setattr(self, k, v)
        self.key = self.sha256 or f"u{self.id}"
        ext, mime, mt, cr = (self.ext or "").lower(), self.mime or "", self.mac_type or "", self.mac_creator or ""
        self.is_pdf = (ext == ".pdf" or mime == "application/pdf" or mt == "PDF ") and not cr.startswith("ART")
        self.is_word = not self.is_pdf and (ext in TEXT_EXT or (cr == "MSWD" and mt in WORD_MAC))
        self.kind = "PDF" if self.is_pdf else "Word/RTF" if self.is_word else (ext.lstrip(".") or mt.strip() or "?")

    def folders(self):
        return " ".join(Path(self.path or "").parts[-3:-1])


def all_files(db):
    return [File(r) for r in db.execute(f"SELECT {', '.join(FILE_COLS)} FROM unique_files")]


def read_files(db, cachedb, texts, files, search, a):
    """Read what hasn't been read yet, a.jobs at a time; results go into texts and the cache."""
    ocr = not a.no_ocr and shutil.which("tesseract") is not None
    if not a.no_ocr and not ocr:
        log("tesseract isn't installed: scans won't be read (brew install tesseract)")
    for tool in ("pdfinfo", "pdftotext", "pdftoppm"):
        if not shutil.which(tool):
            sys.exit(f"{tool} isn't installed: brew install poppler")

    def wanted(f):
        named = bool(search.names and search.names.search(f.name or ""))
        if not (f.is_pdf or f.is_word):
            return False
        got = texts.get(f.key)
        if got is None:
            return True
        if got["text"] is not None:
            return False
        if got["how"].startswith("no text") and not ocr:         # still can't OCR it
            return False
        return f.is_word or named or got["pages"] is None or got["pages"] >= search.min_pages

    todo = [f for f in files if wanted(f)]
    pdf_n = sum(f.is_pdf for f in files)
    log(f"{pdf_n:,} distinct PDFs and {sum(f.is_word for f in files):,} Word/RTF documents; "
        f"{len(todo):,} to read ({len(texts):,} read before)")
    if not todo:
        return 0
    roots = ([str(CONFIG.derivatives_dir)] if CONFIG.derivatives_dir else []) + list(CONFIG.app.get("derivatives_roots") or [])
    copies = pdf_copies(db, roots) if any(f.is_word for f in todo) else {}
    deadline = time.time() + a.max_minutes * 60 if a.max_minutes else float("inf")
    unplugged = done = 0
    last = time.time()

    def job(f, path):
        try:
            if f.is_pdf:
                named = bool(search.names and search.names.search(f.name or ""))
                return read_pdf(path, search.min_pages, named, ocr)
            return read_word(path, copies.get(f.sha256))
        except subprocess.TimeoutExpired:
            return {"pages": None, "text": "", "how": "unreadable: took too long"}
        except Exception as e:
            return {"pages": None, "text": "", "how": f"unreadable: {type(e).__name__}"}

    with cf.ThreadPoolExecutor(a.jobs) as pool:
        running = {}
        it = iter(todo)
        while True:
            while len(running) < a.jobs * 2 and time.time() < deadline:
                f = next(it, None)
                if f is None:
                    break
                path = a_copy(db, f.id, CONFIG.volumes_root)
                if path is None:
                    unplugged += 1
                    continue
                running[pool.submit(job, f, path)] = f
            if not running:
                break
            finished, _ = cf.wait(running, return_when=cf.FIRST_COMPLETED)
            for fut in finished:
                f = running.pop(fut)
                res = fut.result()
                if res["text"]:
                    res["text"] = res["text"][:TEXT_LIMIT]
                texts[f.key] = res
                cachedb.execute("INSERT OR REPLACE INTO doc_text VALUES (?,?,?,?,?)",
                                (f.key, res["pages"], res["text"], res["how"], time.strftime("%Y-%m-%d %H:%M")))
                done += 1
            if time.time() - last > 30:
                last = time.time()
                cachedb.commit()
                log(f"  read {done:,} of {len(todo):,}")
    cachedb.commit()
    if time.time() >= deadline:
        log("time's up: rerun to read the rest")
    if unplugged:
        log(f"{unplugged:,} files are only on drives that aren't plugged in: plug them in and rerun")
    return unplugged


# ---------- matching ----------

def match(files, texts, search, top):
    """-> (matches {target id: [(score, found, missing, File)]}, best {file key: (score, target id)},
    name_only [File])."""
    matches = {t.id: [] for t in search.targets}
    best = {}
    name_only = []
    for f in files:
        named = bool(search.names and search.names.search(f.name or ""))
        got = texts.get(f.key)
        text = got["text"] if got else None
        if not text and not named:
            continue
        if not text:
            name_only.append(f)
        doc = Doc(f"{text or ''}\n{f.name}\n{f.folders()}")
        for t in search.targets:
            s, found, missing = score(t, doc)
            if s >= MIN_SCORE:
                matches[t.id].append((s, found, missing, f))
            if s > best.get(f.key, (0, ""))[0]:            # any score: names alone score low
                best[f.key] = (s, t.id)
    for tid in matches:
        matches[tid].sort(key=lambda m: (-m[0], -(m[3].size_bytes or 0)))
        matches[tid] = matches[tid][:top]
    return matches, best, name_only


# ---------- the spreadsheet ----------

def is_yes(v):
    return str(v or "").strip().lower() in {"yes", "y", "x", "1", "true"}


def old_marks(path):
    """(target id, file key) -> Take, from an earlier spreadsheet, so a rerun keeps JB's marks."""
    if not path.exists():
        return {}
    from openpyxl import load_workbook
    marks = {}
    wb = load_workbook(path, read_only=True)
    for sheet in ("Matches", "By name only"):
        if sheet not in wb.sheetnames:
            continue
        rows = wb[sheet].iter_rows(values_only=True)
        head = list(next(rows, []))
        if not {"Target", "Take", "File key"} <= set(head):
            continue
        it, tk, fk = head.index("Target"), head.index("Take"), head.index("File key")
        for r in rows:
            if r[fk]:
                marks[(r[it] or "", r[fk])] = r[tk] or ""
    return marks


def snippet(text, n=300):
    return " ".join((text or "").split())[:n]


def locations(db, uid, n=10):
    rows = db.execute("SELECT archive, path FROM files WHERE unique_id=? ORDER BY archive, path", (uid,)).fetchall()
    more = f"; ... {len(rows) - n} more" if len(rows) > n else ""
    return "; ".join(f"{a}: {p}" for a, p in rows[:n]) + more


def write_xlsx(path, db, search, files, texts, matches, best, name_only, marks):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter as col

    F = "Arial"
    plain, bold = Font(name=F, size=10), Font(name=F, size=10, bold=True)
    head_fill = PatternFill("solid", fgColor="DDE4EE")
    take_fill = PatternFill("solid", fgColor="E2F0D9")
    thin = Side(style="thin", color="999999")
    wrap = Alignment(wrap_text=True, vertical="top")
    wb = Workbook()

    def sheet(ws, heads, widths):
        for i, (h, w) in enumerate(zip(heads, widths), 1):
            c = ws.cell(1, i, h)
            c.font, c.fill, c.border = bold, head_fill, Border(bottom=thin)
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            ws.column_dimensions[col(i)].width = w
        ws.freeze_panes = "B2"

    def put(ws, r, vals, fill=None):
        for i, v in enumerate(vals, 1):
            c = ws.cell(r, i, v)
            c.font = plain
            c.alignment = wrap
            if fill:
                c.fill = fill
            if isinstance(v, int) and not isinstance(v, bool):
                c.number_format = "#,##0"

    # Targets
    ws = wb.active
    ws.title = "Targets"
    sheet(ws, ["Target", "Title page", "Keys", "Candidates", "Marked Take", "Best score", "Notes"],
          [28, 60, 30, 11, 11, 10, 40])
    taken = {}
    for r, t in enumerate(search.targets, 2):
        ms = matches[t.id]
        n_take = 0
        for s, found, missing, f in ms:
            mark = marks.get((t.id, f.key))
            if is_yes(mark if mark is not None else ("yes" if not missing and s >= TAKE_SCORE else "")):
                n_take += 1
        taken[t.id] = n_take
        put(ws, r, [t.id, t.title, "; ".join(t.keys), len(ms), n_take, ms[0][0] if ms else None, t.notes])
    note_r = len(search.targets) + 3
    ws.cell(note_r, 1, f"Search {search.name} ({search.path.name}), {time.strftime('%Y-%m-%d %H:%M')}. "
                       f"Score: 0.6 x share of the title's words + 0.4 x share of the keys, in the file's "
                       f"first pages, name and folder. Take marks are kept when the search is run again; "
                       f"then: python3 scripts/find_documents.py {search.name} --collect").font = Font(name=F, size=9, italic=True)

    # Matches
    ws = wb.create_sheet("Matches")
    heads = ["Target", "Rank", "Score", "Keys missing", "Take", "Name", "Kind", "Pages", "Size (bytes)",
             "Modified", "Copies", "Read by", "First page", "Where (archive drive: path)",
             "Its best target, if another", "Category", "File key"]
    sheet(ws, heads, [26, 6, 7, 22, 6, 36, 9, 7, 13, 11, 7, 12, 60, 70, 24, 12, 18])
    r = 2
    for t in search.targets:
        if not matches[t.id]:
            put(ws, r, [t.id, None, None, "nothing found"])
            r += 1
            continue
        for rank, (s, found, missing, f) in enumerate(matches[t.id], 1):
            got = texts.get(f.key) or {}
            mark = marks.get((t.id, f.key))
            take = mark if mark is not None else ("yes" if not missing and s >= TAKE_SCORE else "")
            other = best.get(f.key, (0, t.id))[1]
            put(ws, r, [t.id, rank, s, "; ".join(missing), take, f.name, f.kind, got.get("pages"), f.size_bytes,
                        (f.modified or "")[:10], f.copies, got.get("how", "name only"), snippet(got.get("text")),
                        locations(db, f.id), other if other != t.id else "", f.category or "", f.key],
                take_fill if is_yes(take) else None)
            r += 1
    ws.auto_filter.ref = f"A1:{col(len(heads))}{max(r - 1, 1)}"

    # By name only
    ws = wb.create_sheet("By name only")
    heads = ["Target", "Score", "Take", "Name", "Kind", "Size (bytes)", "Modified", "Copies", "Why not read",
             "Where (archive drive: path)", "Category", "File key"]
    sheet(ws, heads, [26, 7, 6, 40, 10, 13, 11, 7, 24, 80, 12, 18])
    for r, f in enumerate(sorted(name_only, key=lambda f: (f.name or "").lower()), 2):
        s, tid = best.get(f.key, (None, ""))
        got = texts.get(f.key)
        why = got["how"] if got else ("its kind isn't read" if not (f.is_pdf or f.is_word) else "drive not plugged in")
        mark = marks.get((tid, f.key), "")
        put(ws, r, [tid, s, mark, f.name, f.kind, f.size_bytes, (f.modified or "")[:10], f.copies, why,
                    locations(db, f.id), f.category or "", f.key], take_fill if is_yes(mark) else None)
    ws.auto_filter.ref = f"A1:{col(len(heads))}{max(len(name_only) + 1, 1)}"

    # Other long PDFs
    ws = wb.create_sheet("Other long PDFs")
    listed = {m[3].key for ms in matches.values() for m in ms}
    others = [f for f in files if f.is_pdf and f.key not in listed and (texts.get(f.key) or {}).get("pages")
              and texts[f.key]["pages"] >= search.min_pages]
    others.sort(key=lambda f: -(f.size_bytes or 0))
    heads = ["Name", "Pages", "Size (bytes)", "Modified", "Copies", "Read by", "First page",
             "Where (archive drive: path)", "Category", "File key"]
    sheet(ws, heads, [40, 7, 13, 11, 7, 12, 70, 80, 12, 18])
    for r, f in enumerate(others, 2):
        got = texts[f.key]
        put(ws, r, [f.name, got["pages"], f.size_bytes, (f.modified or "")[:10], f.copies, got["how"],
                    snippet(got["text"], 200), locations(db, f.id, 3), f.category or "", f.key])
    ws.auto_filter.ref = f"A1:{col(len(heads))}{max(len(others) + 1, 1)}"
    wb.save(path)
    return taken, len(others)


# ---------- collecting ----------

def safe(name):
    return re.sub(r'[/:\x00-\x1f]', "_", name).strip() or "unnamed"


def collect(db, search, xlsx, dest):
    """Copy every file marked Take in the spreadsheet into dest."""
    if not xlsx.exists():
        sys.exit(f"No {xlsx}: run the search first (python3 scripts/find_documents.py {search.name})")
    from openpyxl import load_workbook
    wb = load_workbook(xlsx, read_only=True)
    wanted = []
    for sheet in ("Matches", "By name only"):
        if sheet not in wb.sheetnames:
            continue
        rows = wb[sheet].iter_rows(values_only=True)
        head = list(next(rows, []))
        it, tk, fk, nk = (head.index(h) for h in ("Target", "Take", "File key", "Name"))
        wanted += [(r[it] or "", r[fk], r[nk]) for r in rows if r[fk] and is_yes(r[tk])]
    if not wanted:
        sys.exit(f"Nothing is marked Take in {xlsx}")
    dest.mkdir(parents=True, exist_ok=True)
    manifest, copied, there, missing = [], 0, 0, []
    for tid, key, name in wanted:
        if key.startswith("u") and key[1:].isdigit():
            row = db.execute("SELECT id, name FROM unique_files WHERE id=?", (int(key[1:]),)).fetchone()
        else:
            row = db.execute("SELECT id, name FROM unique_files WHERE sha256=?", (key,)).fetchone()
        src = a_copy(db, row[0], CONFIG.volumes_root) if row else None
        if src is None:
            missing.append(f"{tid}: {name}" + ("" if row else " (no longer in the catalogue)"))
            continue
        out = dest / safe(f"{tid} - {row[1]}" if tid else row[1])
        if out.exists() and out.stat().st_size != src.stat().st_size:
            out = out.with_name(f"{out.stem} ({key[:8]}){out.suffix}")
        if out.exists() and out.stat().st_size == src.stat().st_size:
            there += 1
        else:
            shutil.copy2(src, out)
            copied += 1
        manifest.append((tid, out.name, str(src), key))
    with open(dest / "manifest.csv", "w", newline="", encoding="utf-8-sig") as fh:
        w = csv.writer(fh)
        w.writerow(["Target", "File here", "Copied from", "File key"])
        w.writerows(manifest)
    print(f"{dest}: copied {copied:,}, {there:,} already there; manifest.csv lists them")
    if missing:
        print(f"{len(missing):,} not copied, because no drive holding them is plugged in:")
        for m in missing:
            print(f"  {m}")


# ---------- main ----------

def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("search", nargs="?", help="a search in rules/searches/ (e.g. catalogs), or a CSV file")
    ap.add_argument("--list", action="store_true", help="list the searches")
    ap.add_argument("--collect", action="store_true", help="copy the files marked Take into the found folder")
    ap.add_argument("--found-dir", help="where the spreadsheet and the copies go (default ~/jva_found)")
    ap.add_argument("--top", type=int, default=5, help="candidates listed per document (default 5)")
    ap.add_argument("--min-pages", type=int, help="override the search's min_pages")
    ap.add_argument("--no-ocr", action="store_true", help="don't OCR scans")
    ap.add_argument("--jobs", type=int, default=4, help="files read at once (default 4)")
    ap.add_argument("--max-minutes", type=float, default=0, help="stop reading after this long (0: no limit)")
    ap.add_argument("--catalogue", default=str(CONFIG.catalogue))
    a = ap.parse_args()
    if a.list:
        for p in sorted(SEARCH_DIR.glob("*.csv")):
            s = load_search(p)
            print(f"{s.name:20} {len(s.targets):3} documents  ({p})")
        return
    if not a.search:
        ap.error("which search? (--list shows them)")
    if importlib.util.find_spec("openpyxl") is None:
        sys.exit("This needs openpyxl:  python3 -m pip install openpyxl")
    search = load_search(a.search)
    if a.min_pages is not None:
        search.min_pages = a.min_pages
    found = Path(a.found_dir or CONFIG.raw.get("found_dir") or "~/jva_found").expanduser()
    found.mkdir(parents=True, exist_ok=True)
    xlsx = found / f"{search.name}.xlsx"
    db = connect(a.catalogue, readonly=True)
    if a.collect:
        collect(db, search, xlsx, found / search.name)
        return
    cachedb = open_cache(CONFIG.data_dir / "document_text.sqlite")
    texts = cached(cachedb)
    files = all_files(db)
    try:
        read_files(db, cachedb, texts, files, search, a)
    finally:
        cachedb.commit()
        cachedb.close()
    matches, best, name_only = match(files, texts, search, a.top)
    marks = old_marks(xlsx)
    taken, others = write_xlsx(xlsx, db, search, files, texts, matches, best, name_only, marks)
    db.close()
    found_n = sum(1 for t in search.targets if matches[t.id])
    print(f"Wrote {xlsx}: candidates for {found_n} of {len(search.targets)} documents, "
          f"{sum(1 for v in taken.values() if v)} with a likely match marked Take; "
          f"{len(name_only):,} by name only; {others:,} other long PDFs")
    for t in search.targets:
        ms = matches[t.id]
        print(f"  {t.id:40} {'nothing found' if not ms else f'best {ms[0][0]:.2f}  {ms[0][3].name}'}")


if __name__ == "__main__":
    main()
